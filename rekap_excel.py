"""Impor rekap pesanan harian dari Excel.

Format (sama dengan REKAP_PEMBELIAN.xlsx): baris 1 berisi "TANGGAL"; baris 2 berisi tanggal di kolom kiri tiap pasangan
(mis. 20, 21, 22); baris 3 "NAMA | PESANAN"; baris berikutnya satu baris per item (nama pelanggan + nama item).
Pesanan yang sama diulang = jumlah bertambah (mis. NASI dua baris = 2 porsi).
Tanggal yang sudah ada di aplikasi DITIMPA seluruhnya (tidak ditambah), jadi data tidak dobel.
"""
from __future__ import annotations

import io
import re
from datetime import date, datetime
from difflib import SequenceMatcher

import logic
import orders

ALIAS = {"KRUPUK": "KERUPUK", "KERUPUK": "KERUPUK", "FREZE": "FREEZE", "TELOR": "TELUR", "INDOMIE": "INDOMIE"}
ABAIKAN = {"DAN", "DENGAN", "DG", ""}
BATAS_BARIS_KOSONG = 200      # berhenti membaca setelah sekian baris kosong berturut-turut (abaikan sampah jauh di bawah)
BATAS_BARIS = 20000


def _token(teks: str) -> list[str]:
    t = re.sub(r"[-&+/(),.]", " ", str(teks).upper())
    t = re.sub(r"\bKUKU\s+BIMA\b", "KUKUBIMA", t)
    out = []
    for w in t.split():
        w = ALIAS.get(w, w)
        if w not in ABAIKAN:
            out.append(w)
    return out


def _bersih_nama(n) -> str:
    return " ".join(str(n or "").strip().upper().split())


def _ke_tanggal(v, bulan: int, tahun: int) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        h = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    if 1 <= h <= 31:
        try:
            return date(tahun, bulan, h)
        except ValueError:
            return None
    return None


# ------------------------------------------------------------------ baca berkas
def baca(data: bytes, bulan: int, tahun: int) -> tuple[dict[date, list[tuple[str, str]]], dict]:
    """Kembalikan ({tanggal: [(nama, teks_item), ...]}, info). Info: sheet yang dipakai, baris tanpa nama."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = None
    for ws in wb.worksheets:
        kepala = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if kepala and str(kepala[0] or "").strip().upper().startswith("TANGGAL"):
            sheet = ws
            break
    if sheet is None:
        raise ValueError("Tidak menemukan sheet berformat rekap (sel A1 harus bertuliskan TANGGAL).")

    baris = sheet.iter_rows(min_row=2, max_row=BATAS_BARIS, values_only=True)
    baris_tgl = next(baris, ())
    kolom: list[tuple[date, int]] = []
    for i, v in enumerate(baris_tgl):
        d = _ke_tanggal(v, bulan, tahun) if v not in (None, "") else None
        if d:
            kolom.append((d, i))
    if not kolom:
        raise ValueError("Baris 2 harus berisi tanggal (mis. 20, 21, 22) di kolom nama tiap pasangan.")
    next(baris, None)   # baris judul NAMA | PESANAN

    hasil: dict[date, list[tuple[str, str]]] = {d: [] for d, _ in kolom}
    tanpa_nama = 0
    kosong = 0
    for r in baris:
        if not any(c not in (None, "") for c in r):
            kosong += 1
            if kosong >= BATAS_BARIS_KOSONG:
                break
            continue
        kosong = 0
        for d, i in kolom:
            nama = _bersih_nama(r[i]) if i < len(r) else ""
            teks = str(r[i + 1]).strip() if i + 1 < len(r) and r[i + 1] not in (None, "") else ""
            if not nama and not teks:
                continue
            if not nama or not teks:
                tanpa_nama += 1
                continue
            hasil[d].append((nama, teks))
    wb.close()
    return {d: v for d, v in hasil.items() if v}, {"sheet": sheet.title, "tanpa_nama": tanpa_nama}


# ------------------------------------------------------------------ katalog harga & pencocokan
def katalog(store, tgl: date) -> dict[str, dict]:
    """Semua nama item yang dikenal untuk tanggal tsb. Prioritas harga: menu harian tanggal itu > minuman > menu master."""
    kat: dict[str, dict] = {}
    for r in store.all("Menu").itertuples():
        if r.nama:
            kat[r.nama] = {"kategori": r.kategori or "Lain-lain", "harga": orders._num(r.harga), "hpp": orders._num(r.hpp)}
    for r in store.all("Minuman").itertuples():
        nm = " ".join(f"{r.nama} {r.varian}".split())
        if nm:
            kat[nm] = {"kategori": "Minuman", "harga": orders._num(r.harga), "hpp": orders._num(r.hpp)}
    for r in orders.menu_hari(store, tgl).itertuples():
        kat[r.nama] = {"kategori": r.kategori, "harga": orders._num(r.harga), "hpp": orders._num(r.hpp)}
    return kat


def cocokkan(teks: str, daftar: list[str]) -> str | None:
    """Cari nama item yang paling mirip dengan teks dari Excel; None bila tidak yakin."""
    q = set(_token(teks))
    if not q:
        return None
    terbaik, skor_terbaik = None, (-1, -1.0)
    for nama in daftar:
        c = set(_token(nama))
        if not c:
            continue
        inter = len(q & c)
        contain = inter / len(q)
        jacc = inter / len(q | c)
        if contain < 1.0 and jacc < 0.75:
            continue
        skor = (1 if contain >= 1.0 else 0, jacc)
        if skor > skor_terbaik:
            terbaik, skor_terbaik = nama, skor
    return terbaik


_QTY_DEPAN = re.compile(r"^\s*(\d+)\s*[xX×]?\s+(.+)$")
_QTY_BELAKANG = re.compile(r"^(.+?)\s*[xX×]\s*(\d+)\s*$")
_QTY_ANGKA = re.compile(r"^(.+?)\s+(\d+)\s*$")


def _urai(teks: str, daftar: list[str]) -> tuple[str | None, int, str]:
    """(nama item cocok | None, jumlah, teks kunci untuk laporan). Mengenali '2 NASI', 'NASI x2', 'NASI 2'."""
    t = str(teks).strip()
    cocok = cocokkan(t, daftar)
    if cocok:
        return cocok, 1, t.upper()
    for pola, urut in ((_QTY_DEPAN, (2, 1)), (_QTY_BELAKANG, (1, 2)), (_QTY_ANGKA, (1, 2))):
        m = pola.match(t)
        if m:
            isi, qty = m.group(urut[0]), int(m.group(urut[1]))
            c2 = cocokkan(isi, daftar)
            if c2 and qty > 0:
                return c2, qty, isi.upper()
    return None, 1, t.upper()


# ------------------------------------------------------------------ rencana impor
def rencana(store, data: dict[date, list[tuple[str, str]]], peta_item: dict[str, str | None] | None = None,
            peta_nama: dict[str, str | None] | None = None) -> dict:
    """Susun pesanan dari hasil baca() tanpa menulis apa pun.
    peta_item: {teks Excel -> nama item | None(lewati)} untuk item yang tidak dikenali.
    peta_nama: {nama Excel -> nama akun | None(buat akun baru)} untuk nama yang belum punya akun."""
    peta_item = peta_item or {}
    peta_nama = peta_nama or {}
    akun = store.all("Pelanggan")
    akun = akun[(akun["aktif"] != "0") & (akun["role"] != "admin")]
    dikenal = {n.upper(): n for n in akun["nama"]}

    nama_baru: dict[str, int] = {}
    tak_dikenal: dict[str, int] = {}
    per_tgl: dict[date, dict[str, dict[str, dict]]] = {}
    for tgl, rows in sorted(data.items()):
        kat = katalog(store, tgl)
        daftar = list(kat)
        pesanan: dict[str, dict[str, dict]] = {}
        for nama_x, teks in rows:
            nama = dikenal.get(nama_x)
            if nama is None:
                pilih = peta_nama.get(nama_x)
                nama = dikenal.get(str(pilih).upper()) if pilih else nama_x
                if nama == nama_x:
                    nama_baru[nama_x] = nama_baru.get(nama_x, 0) + 1
            item, qty, kunci = _urai(teks, daftar)
            if item is None:
                alih = peta_item.get(kunci)
                if alih in kat:
                    item, qty = alih, 1
                else:
                    tak_dikenal[kunci] = tak_dikenal.get(kunci, 0) + 1
                    continue
            entri = pesanan.setdefault(nama, {}).setdefault(item, {
                "kategori": kat[item]["kategori"], "item": item, "qty": 0, "harga": kat[item]["harga"], "hpp": kat[item]["hpp"]})
            entri["qty"] += qty
        per_tgl[tgl] = pesanan
    return {"tanggal": per_tgl, "nama_baru": nama_baru, "tak_dikenal": tak_dikenal}


def ringkas(store, rc: dict) -> list[dict]:
    """Baris ringkasan per tanggal untuk ditampilkan sebelum impor."""
    out = []
    for tgl, psn in rc["tanggal"].items():
        ada = orders.pesanan_tgl(store, tgl)
        total = sum(i["qty"] * i["harga"] for its in psn.values() for i in its.values())
        out.append({"tanggal": tgl, "pelanggan": len(psn), "total": total, "ada_sekarang": ada["pelanggan"].nunique()})
    return out


def impor(store, rc: dict, buat_akun: bool, sandi_awal: str, oleh: str = "Impor Excel") -> dict:
    """Tulis ke database. Tiap tanggal pada rencana DITIMPA seluruhnya. Status lunas dipertahankan bila total pesanan sama."""
    sekarang = logic.now_wib().isoformat(timespec="seconds")
    akun_ada = {n.upper() for n in store.all("Pelanggan")["nama"]}
    baru = []
    if buat_akun:
        for nm in rc["nama_baru"]:
            if nm.upper() not in akun_ada:
                akun_ada.add(nm.upper())
                baru.append({"nama": nm, "sandi_hash": logic.hash_sandi(sandi_awal), "role": "pelanggan", "aktif": "1", "dibuat": sekarang})
        store.append("Pelanggan", baru)
    kasbon_lama = store.all("Kasbon")
    lunas_lama = {}
    for oid, g in kasbon_lama[kasbon_lama["status_bayar"] == "Lunas"].groupby("order_id"):
        lunas_lama[oid] = float(g["subtotal"].astype(float).sum())

    jml_pesanan = 0
    dilewati = 0
    tanggal_tulis = []
    for tgl, psn in rc["tanggal"].items():
        rows = []
        for nama, its in psn.items():
            if not buat_akun and nama.upper() not in akun_ada:
                dilewati += 1
                continue
            items = list(its.values())
            baris = orders._baris(tgl, nama, items, "", "Disetujui", f"{tgl.isoformat()}T08:00:00+07:00", oleh)
            if lunas_lama.get(orders.order_id(tgl, nama)) == float(sum(i["qty"] * i["harga"] for i in items)):
                for r in baris:
                    r["status_bayar"] = "Lunas"
            rows += baris
            jml_pesanan += 1
        store.replace_where("Kasbon", "tanggal", tgl.isoformat(), rows)       # menimpa: tidak ditambah
        store.replace_where("Revisi", "tanggal", tgl.isoformat(), [])
        orders.simpan_record(store, tgl, tunggu=True)
        tanggal_tulis.append(tgl.isoformat())
    return {"akun_baru": len(baru), "pesanan": jml_pesanan, "dilewati": dilewati, "tanggal": tanggal_tulis}
